r"""Registro común de modelos y tablero comparativo (Excel).

Cada modelo, venga del notebook que venga (probit, GLMM, boosting, ensambles), se
registra con:

* un **puntaje fuera de pliegue** (OOF) para cada hogar de train, calculado con los
  MISMOS 5 pliegues (`00_data/processed/pliegues.csv`): comparaciones pareadas;
* un **puntaje de test** (modelo re-estimado con todo train);
* metadatos: familia, descripción, hiperparámetros, número de variables y tiempo.

El puntaje está orientado (mayor ⇒ más probable pobre). Si es una probabilidad, se
reportan también métricas probabilísticas (Brier, log-loss, ECE).

`evaluar_todos` recalcula las métricas de TODOS los modelos con el mismo protocolo:

* **F1_anidado** (principal y honesto): el umbral de cada pliegue se elige con los otros cuatro.
* **F1_anidado_sinBOG**: lo mismo, pero sólo con hogares fuera de Bogotá (el test no tiene Bogotá).
* F1 OOF con el umbral óptimo (optimista) e IC95 bootstrap.
* Umbral de envío: el F1-óptimo OOF sin Bogotá (se aplica al puntaje de test).

Archivos: `00_data/processed/modelos/<slug>.parquet` (id, oof), `<slug>_test.parquet`
(id, score) y `<slug>.json` (metadatos). El Excel se escribe en
`reports/resultados_modelos.xlsx`, y los CSV de envío (NO se suben) en `02_outputs/predictions/`.
"""
from __future__ import annotations

import json
import re
import time
from pathlib import Path

import numpy as np
import pandas as pd

from . import metrics as M
from .cv import f1_con_umbral_anidado, bootstrap_ic


def _slug(nombre: str) -> str:
    s = re.sub(r"[^0-9A-Za-z]+", "_", nombre).strip("_")
    return s[:80]


def carpeta(R) -> Path:
    d = Path(R.processed) / "modelos"
    d.mkdir(parents=True, exist_ok=True)
    return d


# Reglas del enunciado (§2.1) y de Kaggle: sólo se envían modelos de CLASIFICACIÓN entrenados con
# uno de estos siete algoritmos. Las familias antiguas de los notebooks 03–05 se mapean aquí.
FAMILIAS_PERMITIDAS = {
    "Linear Probability Model": "Linear Probability Model", "Logit": "Logit", "Elastic Net": "Elastic Net",
    "CART": "CART", "Random Forest": "Random Forest", "Boosting": "Boosting", "Naive Bayes": "Naive Bayes",
    # nombres de familia usados por el notebook 05 (clasificación directa con algoritmos de la lista)
    "Bosques": "Random Forest", "Lineal penalizado": "Elastic Net", "Boosting-parsimonioso": "Boosting",
}


def permitido(meta: dict) -> tuple[bool, str]:
    """(¿se puede enviar a Kaggle?, motivo). Usa el campo `permitido_envio` de los metadatos si
    existe; si no, lo deduce de la familia (regresiones del ingreso, probit, GLMM y ensambles no)."""
    if meta.get("permitido_envio") is not None:
        return bool(meta["permitido_envio"]), meta.get("motivo_envio", "")
    fam = meta.get("familia", "")
    if fam in FAMILIAS_PERMITIDAS:
        return True, f"clasificación directa con {FAMILIAS_PERMITIDAS[fam]} (algoritmo de la lista del enunciado)"
    if "regresi" in fam.lower():
        return False, "ruta de regresión del ingreso (el enunciado exige la ruta de clasificación)"
    if "nsamble" in fam:
        return False, "ensamble/stacking: no es uno de los siete algoritmos del enunciado"
    return False, f"familia '{fam}' fuera de la lista del enunciado (LPM, Logit, EN, CART, RF, Boosting, NB)"


def registrar(R, nombre: str, oof: np.ndarray, test: np.ndarray | None, ids_train, ids_test=None, *,
              familia: str, descripcion: str = "", n_vars: int | None = None, params: dict | None = None,
              es_probabilidad: bool = True, tiempo_s: float | None = None, notebook: str = "",
              archivo_envio: str | None = None, permitido_envio: bool | None = None,
              motivo_envio: str = "") -> str:
    """Guarda las predicciones y los metadatos de un modelo. Devuelve el slug."""
    d = carpeta(R); slug = _slug(nombre)
    pd.DataFrame({"id": np.asarray(ids_train), "oof": np.asarray(oof, float)}).to_parquet(d / f"{slug}.parquet", index=False)
    if test is not None:
        pd.DataFrame({"id": np.asarray(ids_test), "score": np.asarray(test, float)}).to_parquet(d / f"{slug}_test.parquet", index=False)
    meta = {"nombre": nombre, "slug": slug, "familia": familia, "descripcion": descripcion, "n_vars": n_vars,
            "params": params or {}, "es_probabilidad": bool(es_probabilidad), "tiempo_s": tiempo_s,
            "notebook": notebook, "archivo_envio": archivo_envio, "registrado": time.strftime("%Y-%m-%d %H:%M")}
    if permitido_envio is not None:
        meta["permitido_envio"], meta["motivo_envio"] = bool(permitido_envio), motivo_envio
    (d / f"{slug}.json").write_text(json.dumps(meta, ensure_ascii=False, indent=1, default=float), encoding="utf-8")
    return slug


def cargar(R, slug: str):
    d = carpeta(R)
    meta = json.loads((d / f"{slug}.json").read_text(encoding="utf-8"))
    oof = pd.read_parquet(d / f"{slug}.parquet")
    te = pd.read_parquet(d / f"{slug}_test.parquet") if (d / f"{slug}_test.parquet").exists() else None
    return meta, oof, te


def listar(R) -> list[str]:
    return sorted(p.stem for p in carpeta(R).glob("*.json"))


def matriz_oof(R, base_ids, slugs=None) -> pd.DataFrame:
    """DataFrame (hogares × modelos) con los puntajes OOF alineados a `base_ids`."""
    cols = {}
    for s in slugs or listar(R):
        _, oof, _ = cargar(R, s)
        cols[s] = oof.set_index("id").oof.reindex(base_ids).values
    return pd.DataFrame(cols, index=base_ids)


# ------------------------------------------------------------------------- evaluación
def panel(y, s, folds, mask_sinbog, es_prob=True, B=300) -> dict:
    y = np.asarray(y).astype(int); s = np.asarray(s, float)
    f1n, f1k, umb = f1_con_umbral_anidado(y, s, folds)
    t, f1o = M.umbral_optimo_f1(y, s)
    lo, hi, _ = bootstrap_ic(y, s, t, B=B)
    m = np.asarray(mask_sinbog, bool)
    # anidado sin Bogotá: umbrales elegidos y evaluados sólo con hogares fuera de Bogotá
    yhat = np.zeros_like(y)
    for tr, va in folds:
        tr_m, va_m = tr[m[tr]], va[m[va]]
        tk, _ = M.umbral_optimo_f1(y[tr_m], s[tr_m]); yhat[va_m] = (s[va_m] >= tk).astype(int)
    f1n_sb = M.f1(y[m], yhat[m])
    t_sb, f1o_sb = M.umbral_optimo_f1(y[m], s[m])
    yh = (s >= t_sb).astype(int)
    tp, fp, fn, tn = M.confusion(y[m], yh[m])
    out = {"F1_anidado_sinBOG": f1n_sb, "F1_anidado": f1n, "F1_oof_opt_sinBOG": f1o_sb, "F1_oof_opt": f1o,
           "IC95_F1_inf": lo, "IC95_F1_sup": hi, "F1_de_pliegues": float(f1k.std(ddof=1)),
           "F1_pliego_min": float(f1k.min()), "F1_pliego_max": float(f1k.max()),
           "umbral_envio_sinBOG": t_sb, "precision_sinBOG": tp / max(tp + fp, 1), "recall_sinBOG": tp / max(tp + fn, 1),
           "tasa_predicha_sinBOG": float(yh[m].mean()), "tasa_real_sinBOG": float(y[m].mean()),
           "AUC_ROC": M.auc_roc(y, s), "AUC_PR": M.pr_auc(y, s), "AUC_ROC_sinBOG": M.auc_roc(y[m], s[m]),
           "MCC_sinBOG": M.mcc(y[m], yh[m])}
    if es_prob and np.nanmin(s) >= 0 and np.nanmax(s) <= 1:
        out.update({"Brier": M.brier(y, s), "LogLoss": M.log_loss(y, s), "ECE": M.ece(y, s)})
    for k, f in enumerate(f1k):
        out[f"F1_pliegue_{k}"] = float(f)
    return out


def evaluar_todos(R, y, folds, ids_train, mask_sinbog, slugs=None, B=300, verbose=True) -> pd.DataFrame:
    filas = []
    for s in slugs or listar(R):
        meta, oof, te = cargar(R, s)
        sc = oof.set_index("id").oof.reindex(ids_train).values
        if np.isnan(sc).any():
            if verbose:
                print(f"  {s}: OOF incompleto ({np.isnan(sc).sum()} NaN), se omite"); continue
        r = panel(y, sc, folds, mask_sinbog, es_prob=meta.get("es_probabilidad", True), B=B)
        if te is not None:
            r["tasa_predicha_test"] = float((te.score.values >= r["umbral_envio_sinBOG"]).mean())
        ok, motivo = permitido(meta)
        r["permitido_envio"] = "sí" if ok else "no"
        r["motivo_envio"] = motivo
        fila = {"modelo": meta["nombre"], "familia": meta["familia"], "n_vars": meta.get("n_vars"),
                "notebook": meta.get("notebook", ""), "descripcion": meta.get("descripcion", ""), **r,
                "tiempo_s": meta.get("tiempo_s"), "tiene_test": te is not None, "slug": s,
                "params": json.dumps(meta.get("params", {}), ensure_ascii=False, default=str)[:500]}
        filas.append(fila)
        if verbose:
            print(f"  {meta['nombre'][:60]:60s} F1 anidado sinBOG={r['F1_anidado_sinBOG']:.4f}  AUC={r['AUC_ROC']:.4f}")
    t = pd.DataFrame(filas).sort_values("F1_anidado_sinBOG", ascending=False).reset_index(drop=True)
    t.insert(0, "rank", np.arange(1, len(t) + 1))
    return t


def escribir_envio(R, slug: str, umbral: float, sample: pd.DataFrame, nombre_archivo: str | None = None) -> Path:
    """CSV id,pobre con el umbral dado (NO se sube a Kaggle desde aquí)."""
    meta, _, te = cargar(R, slug)
    if te is None:
        raise ValueError(f"{slug} no tiene puntaje de test")
    sub = sample[["id"]].merge(te, on="id", how="left")
    assert sub.score.notna().all() and len(sub) == len(sample)
    sub["pobre"] = (sub.score >= umbral).astype(int)
    nombre_archivo = nombre_archivo or meta.get("archivo_envio") or f"{slug}.csv"
    p = Path(R.submissions) / nombre_archivo
    sub[["id", "pobre"]].to_csv(p, index=False)
    return p


# ------------------------------------------------------------------------- Excel
def excel(ruta: Path, hojas: dict[str, pd.DataFrame], notas: list[str] | None = None,
          resaltar=("F1_anidado_sinBOG", "F1_anidado", "F1_oof_opt_sinBOG", "F1_oof_opt", "AUC_ROC")):
    """Escribe un libro con formato: encabezados, anchos, paneles congelados, escalas de color."""
    from openpyxl.formatting.rule import ColorScaleRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    ruta = Path(ruta)
    with pd.ExcelWriter(ruta, engine="openpyxl") as xw:
        if notas:
            pd.DataFrame({"Notas": notas}).to_excel(xw, sheet_name="Léeme", index=False)
        for nom, df in hojas.items():
            df.to_excel(xw, sheet_name=nom[:31], index=False)
        wb = xw.book
        azul = PatternFill("solid", fgColor="1F3864")
        for ws in wb.worksheets:
            for c in ws[1]:
                c.font = Font(bold=True, color="FFFFFF"); c.fill = azul
                c.alignment = Alignment(wrap_text=True, vertical="center")
            ws.freeze_panes = "C2" if ws.title != "Léeme" else "A2"
            ws.row_dimensions[1].height = 42
            enc = [c.value for c in ws[1]]
            for j, h in enumerate(enc, start=1):
                L = get_column_letter(j)
                vals = [str(ws.cell(i, j).value) for i in range(2, min(ws.max_row, 60) + 1)]
                w = max([len(str(h)) * 0.9] + [len(v) for v in vals]) if vals else len(str(h))
                ws.column_dimensions[L].width = min(max(9, w + 1), 60 if ws.title == "Léeme" else 45)
                if ws.title == "Léeme":
                    ws.column_dimensions[L].width = 160
                    for i in range(2, ws.max_row + 1):
                        ws.cell(i, j).alignment = Alignment(wrap_text=True, vertical="top")
                    continue
                es_num = any(isinstance(ws.cell(i, j).value, float) for i in range(2, min(ws.max_row, 20) + 1))
                if es_num:
                    for i in range(2, ws.max_row + 1):
                        ws.cell(i, j).number_format = "0.0000"
                if h in resaltar and ws.max_row > 2:
                    ws.conditional_formatting.add(f"{L}2:{L}{ws.max_row}", ColorScaleRule(
                        start_type="min", start_color="F8696B", mid_type="percentile", mid_value=50,
                        mid_color="FFEB84", end_type="max", end_color="63BE7B"))
    return ruta
